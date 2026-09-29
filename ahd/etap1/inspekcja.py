"""Szybka inspekcja pobranych plików: nagłówek, liczba wierszy, spójność plików wieloczęściowych."""

from __future__ import annotations

import csv
import re
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterator

EXCEL = {".xlsx", ".xlsm"}
TEXT = {".csv", ".txt"}

# Sufiks części pliku wieloczęściowego: "_cz1", " cz.2", "-part3", "_p4", " (2)".
PART_RE = re.compile(
    r"(?:(?:[ _\-.]*(?:cz(?:esc|ęść)?|part)|[ _\-.]+p)[ _.\-]*\d+|\s*\(\d+\))$", re.IGNORECASE
)


def _decode(sample: bytes) -> tuple[str, str]:
    for enc in ("utf-8-sig", "cp1250"):
        try:
            return sample.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return sample.decode("latin-1"), "latin-1"


def _inspect_text(path: Path) -> dict[str, Any]:
    with path.open("rb") as fh:
        head = fh.read(64 * 1024)
    text, enc = _decode(head)
    first_line = text.splitlines()[0] if text else ""
    try:
        sep = csv.Sniffer().sniff(first_line, delimiters=";,\t|").delimiter
    except csv.Error:
        sep = "\t" if "\t" in first_line else ";"
    header = next(csv.reader([first_line], delimiter=sep), [])
    # Liczenie znaków końca linii jest szybkie także dla setek tysięcy wierszy.
    newlines, last = 0, b""
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            newlines += chunk.count(b"\n")
            last = chunk[-1:]
    lines = newlines + (1 if last and last != b"\n" else 0)
    return {
        "typ": path.suffix.lower().lstrip("."),
        "kodowanie": enc,
        "separator": {"\t": "TAB"}.get(sep, sep),
        "kolumny": [h.strip() for h in header],
        "wiersze": max(lines - 1, 0),
    }


def _inspect_excel(path: Path) -> dict[str, Any]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        rows = ws.max_row
        if rows is None:  # plik bez znacznika wymiaru – trzeba policzyć
            rows = sum(1 for _ in ws.iter_rows(values_only=True))
        return {
            "typ": path.suffix.lower().lstrip("."),
            "arkusz": ws.title,
            "arkusze": len(wb.sheetnames),
            "kolumny": ["" if h is None else str(h).strip() for h in header_row],
            "wiersze": max(rows - 1, 0),
        }
    finally:
        wb.close()


def inspect_file(path: str | Path) -> dict[str, Any]:
    path = Path(path)
    suffix = path.suffix.lower()
    try:
        if suffix in TEXT:
            return _inspect_text(path)
        if suffix in EXCEL:
            return _inspect_excel(path)
        return {"typ": suffix.lstrip("."), "uwaga": "format nieobsługiwany przez inspekcję"}
    except Exception as exc:  # inspekcja nie może przerwać pobierania
        return {"typ": suffix.lstrip("."), "blad": f"{type(exc).__name__}: {exc}"}


def set_name(filename: str) -> str:
    """Nazwa zestawu, do którego należy plik (bez sufiksu części i rozszerzenia)."""
    stem = Path(filename).stem
    return PART_RE.sub("", stem) or stem


def check_sets(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Kontrola plików wieloczęściowych: te same kolumny we wszystkich częściach."""
    groups: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for rec in records:
        groups[(set_name(rec["plik"]), Path(rec["plik"]).suffix.lower())].append(rec)

    issues = []
    for (name, _suffix), recs in sorted(groups.items()):
        if len(recs) < 2:
            continue
        cols = {tuple(r["inspekcja"].get("kolumny", [])) for r in recs}
        total = sum(r["inspekcja"].get("wiersze", 0) for r in recs)
        issue = {
            "zestaw": name,
            "czesci": [r["plik"] for r in recs],
            "wiersze_razem": total,
            "zgodne_kolumny": len(cols) == 1,
        }
        if len(cols) != 1:
            issue["waga"] = "BLAD"
            issue["opis"] = "Części zestawu mają różne kolumny"
        else:
            issue["waga"] = "OK"
            issue["opis"] = f"{len(recs)} części, identyczny układ {len(next(iter(cols)))} kolumn"
        issues.append(issue)
    return issues


def header_signature(columns: list[str]) -> str:
    """Krótki odcisk układu kolumn – pozwala rozpoznać ten sam typ raportu niezależnie od nazwy pliku."""
    import hashlib

    norm = "|".join(c.strip().lower() for c in columns)
    return hashlib.sha1(norm.encode("utf-8")).hexdigest()[:16]


REPORT_TYPES = [
    (re.compile(r"cji3", re.IGNORECASE), "CJI3"),
    (re.compile(r"zrd[_ -]?kkaj", re.IGNORECASE), "ZRD_KKAJ"),
    (re.compile(r"net[_ -]?inv", re.IGNORECASE), "NET_INV"),
]


def guess_report_type(filename: str) -> str:
    """Wstępne rozpoznanie typu raportu po nazwie pliku (do zastąpienia słownikiem sygnatur kolumn)."""
    for pattern, name in REPORT_TYPES:
        if pattern.search(filename):
            return name
    return "nieznany"


def read_rows(path: str | Path, info: dict[str, Any]) -> Iterator[list[Any]]:
    """Wiersze danych (bez nagłówka) jako listy wartości – do załadowania w postaci surowej."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in TEXT:
        sep = {"TAB": "\t"}.get(info.get("separator", ";"), info.get("separator", ";"))
        with path.open("r", encoding=info.get("kodowanie", "utf-8-sig"), newline="") as fh:
            reader = csv.reader(fh, delimiter=sep)
            next(reader, None)
            for row in reader:
                if any(cell.strip() for cell in row):
                    yield row
    elif suffix in EXCEL:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            rows = wb.worksheets[0].iter_rows(min_row=2, values_only=True)
            for row in rows:
                if any(v is not None and str(v).strip() != "" for v in row):
                    yield list(row)
        finally:
            wb.close()
    else:
        raise ValueError(f"Nieobsługiwany format: {path.name}")
